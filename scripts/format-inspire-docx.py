#!/usr/bin/env python3
"""Reformat Desktop 灵感文案新分类.docx with consistent styles."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path("/Users/hanmacbook/Desktop/灵感文案新分类.docx")
OUT = SRC  # overwrite in place


def set_run_font(run, size=11, bold=False, color=None, name="PingFang SC"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def is_group_header(num: int, name: str, declared: str | None) -> bool:
    if declared is not None and len(name) <= 12:
        return True
    verb_starts = (
        "认识",
        "了解",
        "弄懂",
        "探寻",
        "理清",
        "学习",
        "品",
        "读",
        "感",
        "看",
        "回",
        "发",
        "聆",
        "欣",
        "观",
        "尝",
        "辨",
        "留",
        "区",
        "知",
        "追",
        "梳",
        "破",
        "分",
        "避",
        "搞",
        "推",
        "体",
        "整",
        "摘",
        "记",
        "细",
        "学",
        "习",
    )
    if any(name.startswith(v) for v in verb_starts):
        return False
    return num <= 20 and len(name) <= 8


def parse_groups(lines: list[str]) -> list[dict]:
    header_re = re.compile(r"^(\d{1,2})\s*[·・\.．]\s*(.+?)(?:（(\d+)）|\((\d+)\))?$")
    item_re = re.compile(r"^(\d+)[\.．、]\s*(.+)$")

    groups: list[dict] = []
    cur: dict | None = None

    for line in lines:
        if line.startswith("摸摸鱼") or line.startswith("共 ") or line == "灵感文案新分类":
            continue

        hm = header_re.match(line)
        if hm:
            num = int(hm.group(1))
            name = hm.group(2).strip()
            declared = hm.group(3) or hm.group(4)
            if is_group_header(num, name, declared):
                cur = {"num": num, "name": name, "items": []}
                groups.append(cur)
                continue

        if cur is None:
            continue

        im = item_re.match(line)
        if im:
            cur["items"].append(im.group(2).strip())
        else:
            cur["items"].append(line)

    for i, g in enumerate(groups, 1):
        g["num"] = i
    return groups


def main() -> None:
    doc_in = Document(SRC)
    lines = [p.text.strip() for p in doc_in.paragraphs if p.text.strip()]
    groups = parse_groups(lines)

    print("parsed groups:")
    for g in groups:
        print(f"  {g['num']:02d} · {g['name']}（{len(g['items'])}）")

    out_doc = Document()
    for section in out_doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    title = out_doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("灵感文案新分类"), size=18, bold=True)

    sub = out_doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total = sum(len(g["items"]) for g in groups)
    set_run_font(
        sub.add_run(f"共 {len(groups)} 组 · {total} 条"),
        size=10,
        color=RGBColor(0x78, 0x82, 0x8C),
    )

    out_doc.add_paragraph()

    for g in groups:
        h = out_doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(8)
        set_run_font(
            h.add_run(f"{g['num']:02d} · {g['name']}（{len(g['items'])}）"),
            size=14,
            bold=True,
            color=RGBColor(0x20, 0x24, 0x28),
        )

        for i, text in enumerate(g["items"], 1):
            p = out_doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.35)
            set_run_font(p.add_run(f"{i}. {text}"), size=11)

    out_doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
